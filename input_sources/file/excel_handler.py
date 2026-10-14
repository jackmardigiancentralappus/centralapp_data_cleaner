from openpyxl import load_workbook
from user_input import (ask_number)
from datetime import datetime, date

def cell_to_text(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        if value.time() == datetime.min.time():      # date-only cell
            return value.strftime("%Y-%m-%d")
        return value.isoformat(sep=" ")              # real date+time, keep it
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))                       # 4045551212.0 -> "4045551212"
    return str(value)

def parse_excel_records(file_path):
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise ValueError("The Excel sheet is empty.")

        for number, name in enumerate(header, start=1):
            print(f"{number}: {name}")

        id_index = ask_number("ID column number: ", 1, len(header)) - 1
        value_index = ask_number("Column to clean: ", 1, len(header)) - 1

        return [
            {
                "id": cell_to_text(row[id_index]),
                "original_value": cell_to_text(row[value_index]),
            }
            for row in rows
        ]
    finally:
        workbook.close()